from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "src" / "Features" / "Integration" / "data" / "2026_sem1" / "zi" / "01.09.26"


def all_paragraphs(document: Document):
    for element in document.element.body.iter(qn("w:p")):
        yield Paragraph(element, document)


def replace_whole_paragraph(paragraph: Paragraph, replacement: str) -> None:
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(replacement)
        return
    runs[0].text = replacement
    for run in runs[1:]:
        run.text = ""


def remove_one_trailing_parenthesis(paragraph: Paragraph) -> None:
    for run in reversed(paragraph.runs):
        if run.text:
            if not run.text.endswith(")"):
                raise RuntimeError(f"Expected trailing parenthesis in run: {run.text!r}")
            run.text = run.text[:-1]
            return
    raise RuntimeError("Expected a non-empty run")


def repair_year_one() -> None:
    path = DATA / "Orar_An_I Lic.docx"
    document = Document(path)
    replacements = {
        "S e c u r i t a t e a    c i b e r n e t i c ă   (curs)": "Securitatea cibernetică (curs)",
        "Securitatea  cibernetică (curs)": "Securitatea cibernetică (curs)",
    }
    found = {source: 0 for source in replacements}
    for paragraph in all_paragraphs(document):
        if paragraph.text in replacements:
            source = paragraph.text
            replace_whole_paragraph(paragraph, replacements[source])
            found[source] += 1
    canonical_count = sum(
        paragraph.text == "Securitatea cibernetică (curs)"
        for paragraph in all_paragraphs(document)
    )
    if found == {source: 1 for source in replacements}:
        document.save(path)
    elif found != {source: 0 for source in replacements} or canonical_count != 2:
        raise RuntimeError(
            f"Unexpected cybersecurity occurrence counts: {found}, canonical={canonical_count}"
        )


def repair_year_two() -> None:
    path = DATA / "Orar_An_II Lic.docx"
    document = Document(path)
    replacements = {
        "Progr.C++ (lab,I-imp,II-par)": "Progr.C++ (lab, I-imp, II-par)",
        " DCMIMJ(lab,II-imp,I-par)": "DCMIMJ (lab, II-imp, I-par)",
    }
    found = {source: 0 for source in replacements}
    for paragraph in all_paragraphs(document):
        if paragraph.text in replacements:
            source = paragraph.text
            replace_whole_paragraph(paragraph, replacements[source])
            found[source] += 1
    if found != {source: 1 for source in replacements}:
        raise RuntimeError(f"Unexpected compact course occurrence counts: {found}")
    document.save(path)


def repair_year_two_ui_ux_typo() -> None:
    path = DATA / "Orar_An_II Lic.docx"
    document = Document(path)
    typo = "Designul UI/U X(curs)"
    replacement = "Designul UI/UX (curs)"
    found = 0
    for paragraph in all_paragraphs(document):
        if paragraph.text == typo:
            replace_whole_paragraph(paragraph, replacement)
            found += 1
    canonical_count = sum(
        paragraph.text == replacement
        for paragraph in all_paragraphs(document)
    )
    if found == 1:
        if canonical_count != 3:
            raise RuntimeError(
                f"Unexpected canonical UI/UX count after repair: {canonical_count}"
            )
        document.save(path)
    elif found == 0:
        if canonical_count != 3:
            raise RuntimeError(
                f"Typo not present, but canonical UI/UX count is {canonical_count} instead of 3"
            )
    else:
        raise RuntimeError(f"Unexpected UI/UX typo occurrence count: {found}")


if __name__ == "__main__":
    repair_year_one()
    repair_year_two()
    repair_year_two_ui_ux_typo()
